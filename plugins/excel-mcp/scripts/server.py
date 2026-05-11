from __future__ import annotations

import json
import math
import os
from pathlib import Path
from typing import Any

import pandas as pd
import requests
from mcp.server.fastmcp import FastMCP

try:
    from openpyxl import Workbook, load_workbook
except ImportError as exc:  # pragma: no cover
    raise RuntimeError("openpyxl is required for excel-mcp") from exc


SERVER = FastMCP(
    "excel-mcp",
    instructions=(
        "Excel workbook tooling for governed reporting packs, workbook QA, and "
        "Microsoft Graph workbook access."
    ),
    json_response=True,
)


def _normalize_path(path_text: str) -> Path:
    path = Path(path_text).expanduser()
    if not path.is_absolute():
        path = Path.cwd() / path
    return path.resolve()


def _sheet_title(name: str) -> str:
    cleaned = name.replace("\\", "_").replace("/", "_").replace("*", "_")
    cleaned = cleaned.replace("[", "_").replace("]", "_").replace(":", "_").replace("?", "_")
    return cleaned[:31] or "Sheet1"


def _coerce_scalar(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (pd.Timestamp, pd.Timedelta)):
        return str(value)
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return None
        return round(value, 6)
    if pd.isna(value):
        return None
    return value


def _records(df: pd.DataFrame, max_rows: int | None = None) -> list[dict[str, Any]]:
    if max_rows is not None:
        df = df.head(max_rows)
    return [
        {str(key): _coerce_scalar(value) for key, value in row.items()}
        for row in df.to_dict(orient="records")
    ]


def _load_tabular_data(path_text: str, sheet_name: str | None = None) -> tuple[pd.DataFrame, dict[str, Any]]:
    path = _normalize_path(path_text)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        df = pd.read_csv(path)
        meta = {"source_type": "csv"}
    elif suffix == ".parquet":
        df = pd.read_parquet(path)
        meta = {"source_type": "parquet"}
    elif suffix in {".json", ".jsonl"}:
        df = pd.read_json(path, lines=(suffix == ".jsonl"))
        meta = {"source_type": "json"}
    elif suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        target_sheet = sheet_name or 0
        df = pd.read_excel(path, sheet_name=target_sheet)
        resolved_sheet = sheet_name
        if isinstance(target_sheet, int):
            workbook = load_workbook(path, read_only=True, data_only=False)
            resolved_sheet = workbook.sheetnames[target_sheet]
        meta = {"source_type": "excel", "sheet_name": resolved_sheet or target_sheet}
    else:
        raise ValueError(f"Unsupported file extension: {suffix}")
    return df, meta


def _profile_dataframe(df: pd.DataFrame, max_categories: int = 10) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for column in df.columns:
        series = df[column]
        non_null = series.dropna()
        row: dict[str, Any] = {
            "column": str(column),
            "dtype": str(series.dtype),
            "rows": int(series.shape[0]),
            "nulls": int(series.isna().sum()),
            "null_pct": round(float(series.isna().mean() * 100), 2) if len(series) else 0.0,
            "distinct_non_null": int(non_null.nunique(dropna=True)),
        }
        if pd.api.types.is_numeric_dtype(series):
            row["min"] = _coerce_scalar(non_null.min()) if not non_null.empty else None
            row["max"] = _coerce_scalar(non_null.max()) if not non_null.empty else None
            row["mean"] = _coerce_scalar(non_null.mean()) if not non_null.empty else None
            row["sum"] = _coerce_scalar(non_null.sum()) if not non_null.empty else None
        elif pd.api.types.is_datetime64_any_dtype(series):
            row["min"] = _coerce_scalar(non_null.min()) if not non_null.empty else None
            row["max"] = _coerce_scalar(non_null.max()) if not non_null.empty else None
        else:
            top_values = non_null.astype(str).value_counts().head(max_categories).to_dict()
            row["top_values"] = json.dumps(top_values)
        rows.append(row)
    return pd.DataFrame(rows)


def _qa_checks(df: pd.DataFrame, dimension_columns: list[str], metric_columns: list[str]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = [
        {"check": "row_count", "value": int(len(df)), "status": "info"},
        {"check": "column_count", "value": int(len(df.columns)), "status": "info"},
        {"check": "duplicate_rows", "value": int(df.duplicated().sum()), "status": "warn"},
    ]
    for column in dimension_columns:
        if column in df.columns:
            rows.append(
                {"check": f"missing_{column}", "value": int(df[column].isna().sum()), "status": "warn"}
            )
    for column in metric_columns:
        if column in df.columns and pd.api.types.is_numeric_dtype(df[column]):
            rows.append(
                {
                    "check": f"negative_{column}",
                    "value": int((df[column].fillna(0) < 0).sum()),
                    "status": "warn",
                }
            )
    return pd.DataFrame(rows)


def _summarize_numeric(df: pd.DataFrame, metric_columns: list[str]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = [{"metric": "row_count", "value": int(len(df))}]
    for column in metric_columns:
        if column in df.columns and pd.api.types.is_numeric_dtype(df[column]):
            series = df[column].dropna()
            rows.extend(
                [
                    {"metric": f"{column}_sum", "value": _coerce_scalar(series.sum())},
                    {"metric": f"{column}_avg", "value": _coerce_scalar(series.mean())},
                    {"metric": f"{column}_max", "value": _coerce_scalar(series.max())},
                ]
            )
    return pd.DataFrame(rows)


def _build_reporting_frames(
    df: pd.DataFrame,
    date_column: str | None,
    dimension_columns: list[str],
    metric_columns: list[str],
    top_n: int,
) -> dict[str, pd.DataFrame]:
    numeric_columns = [
        column
        for column in (metric_columns or list(df.select_dtypes(include=["number"]).columns))
        if column in df.columns
    ]
    dimension_columns = [column for column in dimension_columns if column in df.columns]
    frames: dict[str, pd.DataFrame] = {
        "raw_sample": df.head(200),
        "profile": _profile_dataframe(df),
        "kpi_summary": _summarize_numeric(df, numeric_columns),
        "qa_checks": _qa_checks(df, dimension_columns, numeric_columns),
    }
    if date_column and date_column in df.columns:
        dates = pd.to_datetime(df[date_column], errors="coerce")
        trend_df = df.copy()
        trend_df["_period"] = dates.dt.to_period("M").astype("string")
        trend_df = trend_df.dropna(subset=["_period"])
        if numeric_columns:
            frames["monthly_trend"] = (
                trend_df.groupby("_period", dropna=False)[numeric_columns]
                .sum(numeric_only=True)
                .reset_index()
            )
        else:
            frames["monthly_trend"] = trend_df.groupby("_period", dropna=False).size().reset_index(name="row_count")
    for column in dimension_columns[:3]:
        sheet_name = _sheet_title(f"by_{column}")
        if numeric_columns:
            grouped = (
                df.groupby(column, dropna=False)[numeric_columns]
                .sum(numeric_only=True)
                .sort_values(by=numeric_columns[0], ascending=False)
                .head(top_n)
                .reset_index()
            )
        else:
            grouped = (
                df.groupby(column, dropna=False)
                .size()
                .reset_index(name="row_count")
                .sort_values(by="row_count", ascending=False)
                .head(top_n)
            )
        frames[sheet_name] = grouped
    return frames


def _write_rows_to_sheet(workbook_path: Path, sheet_name: str, rows: list[dict[str, Any]], replace: bool) -> dict[str, Any]:
    if workbook_path.exists():
        workbook = load_workbook(workbook_path)
    else:
        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)
    safe_name = _sheet_title(sheet_name)
    if safe_name in workbook.sheetnames and replace:
        workbook.remove(workbook[safe_name])
    elif safe_name in workbook.sheetnames:
        raise ValueError(f"Sheet already exists: {safe_name}")
    worksheet = workbook.create_sheet(title=safe_name)
    if rows:
        columns = list(rows[0].keys())
        worksheet.append(columns)
        for row in rows:
            worksheet.append([row.get(column) for column in columns])
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(workbook_path)
    return {"workbook_path": str(workbook_path), "sheet_name": safe_name, "rows_written": len(rows)}


def _graph_headers() -> dict[str, str]:
    token = os.environ.get("MS_GRAPH_ACCESS_TOKEN", "").strip()
    if not token:
        raise RuntimeError("MS_GRAPH_ACCESS_TOKEN is not set.")
    return {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}


def _graph_workbook_url(drive_id: str, item_id: str, suffix: str) -> str:
    return f"https://graph.microsoft.com/v1.0/drives/{drive_id}/items/{item_id}/workbook/{suffix.lstrip('/')}"


@SERVER.tool()
def excel_inspect_workbook(workbook_path: str, sample_rows: int = 10) -> dict[str, Any]:
    """Inspect a local workbook and return sheet-level metadata."""
    path = _normalize_path(workbook_path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets: list[dict[str, Any]] = []
    for name in workbook.sheetnames:
        worksheet = workbook[name]
        values = list(worksheet.iter_rows(values_only=True, max_row=sample_rows))
        header = [str(value) if value is not None else "" for value in values[0]] if values else []
        sheets.append(
            {
                "sheet_name": name,
                "max_row": int(worksheet.max_row or 0),
                "max_column": int(worksheet.max_column or 0),
                "header": header,
                "sample_rows": [list(row) for row in values[1:sample_rows]],
            }
        )
    return {
        "workbook_path": str(path),
        "sheet_count": len(sheets),
        "sheet_names": workbook.sheetnames,
        "sheets": sheets,
    }


@SERVER.tool()
def excel_read_sheet(workbook_path: str, sheet_name: str, max_rows: int = 100) -> dict[str, Any]:
    """Read a sheet from a local workbook into JSON-ready records."""
    df, meta = _load_tabular_data(workbook_path, sheet_name=sheet_name)
    return {
        "workbook_path": str(_normalize_path(workbook_path)),
        "sheet_name": meta.get("sheet_name", sheet_name),
        "rows_returned": min(len(df), max_rows),
        "columns": [str(column) for column in df.columns],
        "rows": _records(df, max_rows=max_rows),
    }


@SERVER.tool()
def excel_write_sheet(
    workbook_path: str,
    sheet_name: str,
    rows: list[dict[str, Any]],
    replace_sheet: bool = True,
) -> dict[str, Any]:
    """Write JSON-style rows to a local workbook sheet."""
    return _write_rows_to_sheet(_normalize_path(workbook_path), sheet_name, rows, replace_sheet)


@SERVER.tool()
def excel_profile_table(
    source_path: str,
    sheet_name: str | None = None,
    max_categories: int = 10,
) -> dict[str, Any]:
    """Profile a local tabular file or workbook sheet."""
    df, meta = _load_tabular_data(source_path, sheet_name=sheet_name)
    profile = _profile_dataframe(df, max_categories=max_categories)
    return {
        "source_path": str(_normalize_path(source_path)),
        "source_meta": meta,
        "row_count": int(len(df)),
        "column_count": int(len(df.columns)),
        "profile": _records(profile),
    }


@SERVER.tool()
def excel_build_reporting_pack(
    source_path: str,
    output_workbook_path: str,
    date_column: str | None = None,
    dimension_columns: list[str] | None = None,
    metric_columns: list[str] | None = None,
    top_n: int = 10,
) -> dict[str, Any]:
    """Build a reporting workbook from a governed dataset."""
    df, meta = _load_tabular_data(source_path)
    frames = _build_reporting_frames(
        df=df,
        date_column=date_column,
        dimension_columns=dimension_columns or [],
        metric_columns=metric_columns or [],
        top_n=top_n,
    )
    output_path = _normalize_path(output_workbook_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, frame in frames.items():
            frame.to_excel(writer, sheet_name=_sheet_title(sheet_name), index=False)
    return {
        "source_path": str(_normalize_path(source_path)),
        "output_workbook_path": str(output_path),
        "source_meta": meta,
        "row_count": int(len(df)),
        "sheets_created": list(frames.keys()),
    }


@SERVER.tool()
def excel_graph_list_worksheets(drive_id: str, item_id: str) -> dict[str, Any]:
    """List worksheets in a Microsoft 365 workbook using Microsoft Graph."""
    response = requests.get(
        _graph_workbook_url(drive_id, item_id, "worksheets"),
        headers=_graph_headers(),
        timeout=30,
    )
    response.raise_for_status()
    payload = response.json()
    return {
        "drive_id": drive_id,
        "item_id": item_id,
        "worksheets": [
            {"id": item.get("id"), "name": item.get("name"), "position": item.get("position")}
            for item in payload.get("value", [])
        ],
    }


@SERVER.tool()
def excel_graph_read_range(
    drive_id: str,
    item_id: str,
    worksheet_name: str,
    address: str,
) -> dict[str, Any]:
    """Read a worksheet range from Microsoft Graph."""
    response = requests.get(
        _graph_workbook_url(
            drive_id,
            item_id,
            f"worksheets/{worksheet_name}/range(address='{address}')",
        ),
        headers=_graph_headers(),
        timeout=30,
    )
    response.raise_for_status()
    payload = response.json()
    return {
        "drive_id": drive_id,
        "item_id": item_id,
        "worksheet_name": worksheet_name,
        "address": address,
        "values": payload.get("values", []),
        "text": payload.get("text", []),
    }


@SERVER.tool()
def excel_graph_write_range(
    drive_id: str,
    item_id: str,
    worksheet_name: str,
    address: str,
    values: list[list[Any]],
) -> dict[str, Any]:
    """Write values to a worksheet range through Microsoft Graph."""
    response = requests.patch(
        _graph_workbook_url(
            drive_id,
            item_id,
            f"worksheets/{worksheet_name}/range(address='{address}')",
        ),
        headers=_graph_headers(),
        json={"values": values},
        timeout=30,
    )
    response.raise_for_status()
    return {
        "drive_id": drive_id,
        "item_id": item_id,
        "worksheet_name": worksheet_name,
        "address": address,
        "rows_written": len(values),
    }


if __name__ == "__main__":
    SERVER.run()
